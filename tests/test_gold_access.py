"""Gold tables as files that people open: ``ldp gold export``, ``ldp gold duckdb`` and ``ldp gold build --duckdb``.

A gold table is written as a CSV, Parquet, JSONL or Excel file, and the whole gold layer is written as one DuckDB file
with the tables, what each table and column means, a ``_catalog`` table and a ``_metrics`` table. Every reader reads the
current release by default, ``--release`` an older one and ``--live`` the tables as they are now, and the files are
written to a temporary path and renamed, so a failure leaves the file that was there. Everything here is offline. The
gold layers are small ones that the tests build under ``tmp_path``.
"""

from __future__ import annotations

import datetime as dt
import decimal
import gc
import json
import os
import struct
from pathlib import Path

import pyarrow as pa
import pyarrow.csv as pacsv
import pyarrow.parquet as pq
import pytest

from local_data_platform import datasets
from local_data_platform.cli import main
from local_data_platform.etl import run_config
from local_data_platform.exceptions import ConfigError, EngineNotFound
from local_data_platform.format.excel import Excel
from local_data_platform.gold import release as releases
from local_data_platform.gold import serve
from local_data_platform.gold.build import build_gold
from tests.test_gold_build import GOLD_CATALOG, external, gold, silver, table, write
from tests.test_gold_release import more_orders

duckdb = pytest.importorskip("duckdb")
openpyxl = pytest.importorskip("openpyxl")

UTC = dt.timezone.utc
REVENUE = [{"restaurant": "north", "orders": 2, "revenue": 30.5}, {"restaurant": "south", "orders": 1, "revenue": 7.25}]
METRICS = {
    "apiVersion": "ldp/v1", "kind": "Metrics", "identifier": "sales", "who": "finance",
    "semantic_models": [{
        "name": "revenue_model", "table": "gold.gold_revenue", "description": "Revenue of each restaurant.",
        "dimensions": ["restaurant"],
        "measures": [{"name": "revenue_sum", "agg": "sum", "expr": "revenue", "description": "What was sold."},
                     {"name": "order_count", "agg": "sum", "expr": "orders"}]}],
    "metrics": [
        {"name": "revenue", "type": "simple", "description": "What was sold.", "measure": "revenue_sum",
         "format": "currency"},
        {"name": "avg_ticket", "type": "ratio", "description": "Revenue per order.", "numerator": "revenue_sum",
         "denominator": "order_count", "format": "number", "owner": "kitchen"}]}
TYPES_COLUMNS = {
    "id": "The id.", "name": "A name, with quotes.", "day": "A date.", "moment": "A time with a zone.",
    "local_moment": "A time with no zone.", "amount": "A decimal.",
    "tags": "A list.", "info": "A record.", "formula": "=1+1 is text, and not a sum.", "ratio": "A number or none.",
    "big": "An integer that a sheet would round."}
TYPES_SQL = """
SELECT * FROM (VALUES
  (1, 'north', DATE '2026-10-01', TIMESTAMPTZ '2026-10-01 08:30:00+00', TIMESTAMP '2026-10-01 08:30:00',
   12.50::DECIMAL(10, 2), [1, 2], {'x': 1}, '=1+1', NULL::DOUBLE, 123456789012345678::BIGINT),
  (2, 'say "hi", ok', DATE '2026-10-02', TIMESTAMPTZ '2026-10-02 09:00:00+02', TIMESTAMP '2026-10-02 09:00:00',
   7.25::DECIMAL(10, 2), [3], {'x': 2}, 'plain', 1.5, 5::BIGINT)
) AS t(id, name, day, moment, local_moment, amount, tags, info, formula, ratio, big)
"""


@pytest.fixture
def layer(tmp_path) -> Path:
    """A folder with ``gold.gold_revenue``, which has a contract, a metrics file, and release r1."""
    silver(tmp_path)
    gold(tmp_path, "gold_revenue", docs={"owner": "finance", "primary_key": ["restaurant"]}, owner="finance")
    write(tmp_path / "metrics.json", METRICS)
    result = build_gold(tmp_path, release=True)
    assert result.exit_code == 0 and result.release.cut and result.release.release.name == "r1"
    return tmp_path


@pytest.fixture
def checked(tmp_path) -> Path:
    """A folder with one gold table that something else builds, so a build checks it and a second build cuts nothing."""
    external(tmp_path)
    return tmp_path


def cli(capsys, *argv) -> tuple[int, str, str]:
    """Run ``ldp``, and leave out the one warning of a machine whose DuckDB has no iceberg extension."""
    code = main([str(item) for item in argv])
    out, err = capsys.readouterr()
    kept = [line for line in err.splitlines(keepends=True) if "Registering Iceberg tables in memory" not in line]
    return code, out, "".join(kept)


def csv_rows(path: Path) -> list[dict]:
    return pacsv.read_csv(path).to_pylist()


def open_file(path: Path):
    """A read-only connection that loads no extension, as a BI tool with a plain DuckDB has."""
    return duckdb.connect(str(path), read_only=True,
                          config={"autoload_known_extensions": False, "autoinstall_known_extensions": False})


def query(path: Path, sql: str) -> list[tuple]:
    connection = open_file(path)
    try:
        return connection.execute(sql).fetchall()
    finally:
        connection.close()


def leftovers(folder: Path) -> list[str]:
    """What a write leaves behind that is not a finished file."""
    return sorted(item.name for item in folder.iterdir() if item.name.endswith((".tmp", ".wal")))


# --------------------------------------------------------------------------- one table as a file


def test_a_table_is_written_as_csv_parquet_and_jsonl_and_the_format_comes_from_the_end_of_the_name(layer, tmp_path):
    out = tmp_path / "out"

    csv = serve.export_table(layer, "gold_revenue", out / "revenue.csv")
    parquet = serve.export_table(layer, "gold_revenue", out / "revenue.parquet")
    jsonl = serve.export_table(layer, "gold_revenue", out / "revenue.jsonl")

    assert [item.format for item in (csv, parquet, jsonl)] == ["csv", "parquet", "jsonl"]
    assert csv_rows(out / "revenue.csv") == REVENUE
    assert pq.read_table(out / "revenue.parquet").to_pylist() == REVENUE
    assert [json.loads(line) for line in (out / "revenue.jsonl").read_text().splitlines()] == REVENUE
    assert (csv.table, csv.rows, csv.columns) == ("gold.gold_revenue", 2, 3)
    assert leftovers(out) == []


def test_a_name_in_capitals_and_a_format_that_is_given_work_and_an_unknown_one_is_refused(layer, tmp_path):
    serve.export_table(layer, "gold_revenue", tmp_path / "A.CSV")
    serve.export_table(layer, "gold_revenue", tmp_path / "data.txt", format=" Csv ")

    assert csv_rows(tmp_path / "A.CSV") == REVENUE == csv_rows(tmp_path / "data.txt")
    unknown = r"cannot tell the format of .*x\.dat from its name\. End the name with one of \['\.csv'"
    for dest, given, message in (("x.dat", None, unknown),
                                 ("x.csv", "xml", r"unknown export format 'xml'"),
                                 ("x.csv", "xlsx", r"x\.csv ends in \.csv, which is csv, and the format is xlsx")):
        with pytest.raises(ConfigError, match=message):
            serve.export_table(layer, "gold_revenue", tmp_path / dest, format=given)
    assert not (tmp_path / "x.dat").exists() and not (tmp_path / "x.csv").exists()


def test_the_parquet_file_says_which_table_release_and_snapshot_it_was_written_from(layer, tmp_path):
    done = serve.export_table(layer, "gold_revenue", tmp_path / "revenue.parquet")

    stamp = json.loads(pq.read_schema(tmp_path / "revenue.parquet").metadata[b"ldp.gold"])

    assert stamp == {"table": "gold.gold_revenue", "release": "r1", "snapshot_id": done.snapshot_id}
    assert done.snapshot_id == releases.current_release(layer).table("gold.gold_revenue").snapshot_id


def test_a_table_is_read_as_the_current_release_holds_it_and_an_older_release_and_live_are_asked_for(layer, tmp_path):
    more_orders(layer)
    build_gold(layer, release=False)  # gold now has the new restaurant, and r1 is still the current release

    first = serve.export_table(layer, "gold_revenue", tmp_path / "default.csv")
    live = serve.export_table(layer, "gold_revenue", tmp_path / "live.csv", live=True)

    assert csv_rows(tmp_path / "default.csv") == REVENUE
    assert [row["restaurant"] for row in csv_rows(tmp_path / "live.csv")] == ["extra0", "north", "south"]
    assert (first.release, first.rows) == ("r1", 2) and (live.release, live.rows) == (None, 3)
    assert live.snapshot_id == table(layer, "gold_revenue").table().current_snapshot().snapshot_id != first.snapshot_id
    assert first.source.startswith("release r1, cut ") and live.source == "the tables as they are now"

    build_gold(layer, release=True)  # r2
    newest = serve.export_table(layer, "gold_revenue", tmp_path / "r2.csv")
    older = serve.export_table(layer, "gold_revenue", tmp_path / "r1.csv", release="r1")

    assert (newest.release, newest.rows, older.release, older.rows) == ("r2", 3, "r1", 2)
    assert csv_rows(tmp_path / "r1.csv") == REVENUE and len(csv_rows(tmp_path / "r2.csv")) == 3


def test_a_table_that_the_current_release_does_not_hold_is_read_as_it_is_and_the_result_says_why(layer, tmp_path):
    run_config(gold(layer, "gold_newer", docs={"description": "Added after r1."}))

    done = serve.export_table(layer, "gold_newer", tmp_path / "newer.csv")

    assert done.release is None and done.rows == 2
    assert done.source == "the tables as they are now, because release r1 does not hold gold.gold_newer"
    with pytest.raises(ConfigError, match=r"release r1 does not hold gold\.gold_newer, so it cannot be read as of "
                                          r"that release\. The release holds \['gold\.gold_revenue'\]"):
        serve.export_table(layer, "gold_newer", tmp_path / "older.csv", release="r1")
    assert not (tmp_path / "older.csv").exists()


def test_a_folder_with_no_release_yet_is_read_as_it_is_and_says_so(tmp_path):
    silver(tmp_path)
    run_config(gold(tmp_path))

    done = serve.export_table(tmp_path, "gold_revenue", tmp_path / "revenue.csv")

    assert done.release is None and csv_rows(tmp_path / "revenue.csv") == REVENUE
    assert done.source == "the tables as they are now, because no release has been cut yet"


def test_a_release_that_is_not_there_and_a_release_with_live_are_refused(layer, tmp_path):
    with pytest.raises(releases.ReleaseError, match="r9"):
        serve.export_table(layer, "gold_revenue", tmp_path / "x.csv", release="r9")
    with pytest.raises(ConfigError, match="either a release or --live"):
        serve.export_table(layer, "gold_revenue", tmp_path / "x.csv", release="r1", live=True)
    assert not (tmp_path / "x.csv").exists()


def test_a_table_is_found_by_its_name_or_its_bare_name_and_an_unknown_one_lists_the_tables(layer, tmp_path):
    assert serve.export_table(layer, "gold.gold_revenue", tmp_path / "a.csv").table == "gold.gold_revenue"
    assert serve.export_table(layer, "gold_revenue", tmp_path / "b.csv").table == "gold.gold_revenue"

    with pytest.raises(ConfigError, match=r"gold_missing.*gold\.gold_revenue"):
        serve.export_table(layer, "gold_missing", tmp_path / "c.csv")


def test_a_snapshot_that_a_release_names_and_that_is_gone_is_an_error_and_nothing_is_written(
        layer, tmp_path, monkeypatch):
    monkeypatch.setattr(releases, "snapshot_for", lambda name, chosen: 12345 if chosen is not None else None)

    with pytest.raises(releases.ReleaseError, match="names snapshot 12345 of gold.gold_revenue, which no longer"):
        serve.export_table(layer, "gold_revenue", tmp_path / "x.csv")

    assert not (tmp_path / "x.csv").exists()


def test_a_failed_write_leaves_the_file_that_was_there_and_no_temporary_file(layer, tmp_path, monkeypatch):
    target = tmp_path / "revenue.csv"
    serve.export_table(layer, "gold_revenue", target)
    before = target.read_bytes()
    more_orders(layer)
    build_gold(layer, release=True)

    def broken(sink, *args, **kwargs):
        raise OSError("the disk is full")

    monkeypatch.setattr("pyarrow.csv.write_csv", broken)
    with pytest.raises(OSError, match="disk is full"):
        serve.export_table(layer, "gold_revenue", target)

    assert target.read_bytes() == before and leftovers(tmp_path) == []


# --------------------------------------------------------------------------- the same through the commands


def test_ldp_gold_export_prints_what_it_wrote_and_what_it_read(layer, tmp_path, capsys):
    code, out, err = cli(capsys, "gold", "export", layer, "gold_revenue", tmp_path / "revenue.csv")

    assert code == 0 and err == ""
    assert out.startswith(f"wrote {tmp_path / 'revenue.csv'}, 2 rows of gold.gold_revenue, read from release r1, cut ")
    assert out.rstrip().endswith(f"snapshot {releases.current_release(layer).table('gold.gold_revenue').snapshot_id}")
    assert csv_rows(tmp_path / "revenue.csv") == REVENUE


def test_the_json_of_an_export_has_the_table_the_file_the_release_and_the_snapshot(layer, tmp_path, capsys):
    code, out, _ = cli(capsys, "gold", "export", layer, "gold_revenue", tmp_path / "r.jsonl", "--json")

    data = json.loads(out)

    assert code == 0 and data["dest"] == str(tmp_path / "r.jsonl")
    assert (data["table"], data["format"], data["rows"], data["columns"], data["release"]) == (
        "gold.gold_revenue", "jsonl", 2, 3, "r1")
    assert data["snapshot_id"] == releases.current_release(layer).table("gold.gold_revenue").snapshot_id


def test_ldp_gold_export_takes_a_format_a_release_and_live_and_not_a_release_with_live(layer, tmp_path, capsys):
    more_orders(layer)
    build_gold(layer, release=True)  # r2 has the new restaurant, and r1 does not

    code, _, _ = cli(capsys, "gold", "export", layer, "gold_revenue", tmp_path / "x.dat", "--format", "csv", "--live")
    assert code == 0 and len(csv_rows(tmp_path / "x.dat")) == 3
    code, _, _ = cli(capsys, "gold", "export", layer, "gold_revenue", tmp_path / "y.csv", "--release", "r1")
    assert code == 0 and csv_rows(tmp_path / "y.csv") == REVENUE
    code, _, err = cli(capsys, "gold", "export", layer, "gold_revenue", tmp_path / "z.csv", "--release", "r1", "--live")
    assert code == 1 and "not allowed with argument" in err and not (tmp_path / "z.csv").exists()


def test_ldp_gold_export_says_in_one_line_what_is_wrong(layer, tmp_path, capsys):
    code, out, err = cli(capsys, "gold", "export", layer, "gold_revenue", tmp_path / "x.dat")

    assert code == 1 and out == ""
    assert err.startswith("ldp: error: ConfigError: cannot tell the format of ") or "cannot tell the format" in err


# --------------------------------------------------------------------------- an Excel file


@pytest.fixture
def types(tmp_path) -> Path:
    """A folder with ``gold.gold_types``, a table of every kind of value that a sheet has to take care of."""
    silver(tmp_path)
    gold(tmp_path, "gold_types", sql=TYPES_SQL, columns=TYPES_COLUMNS,
         docs={"grain": "one row for each id", "description": "Every kind of value.", "owner": "ops"})
    build_gold(tmp_path, release=True)
    return tmp_path


def test_an_excel_file_has_the_rows_on_a_sheet_of_the_table_and_the_contract_on_a_sheet_called_about(layer, tmp_path):
    done = serve.export_table(layer, "gold_revenue", tmp_path / "revenue.xlsx")

    book = openpyxl.load_workbook(tmp_path / "revenue.xlsx")

    assert done.format == "xlsx" and book.sheetnames == ["gold_revenue", "about"]
    assert [list(row) for row in book["gold_revenue"].iter_rows(values_only=True)] == [
        ["restaurant", "orders", "revenue"], ["north", 2, 30.5], ["south", 1, 7.25]]
    about = [[cell for cell in row if cell is not None] for row in book["about"].iter_rows(values_only=True)]
    snapshot = releases.current_release(layer).table("gold.gold_revenue").snapshot_id
    assert about[:9] == [
        ["Table", "gold.gold_revenue"], ["Description", "The gold_revenue table."], ["Owner", "finance"],
        ["Grain", "one row per restaurant"], ["Primary key", "restaurant"], ["Release", "r1"],
        ["Snapshot id", snapshot if snapshot <= 2 ** 53 else str(snapshot)], ["Built at (UTC)", about[7][1]],
        ["Rows", 2]]
    assert about[9][0] == "Exported at (UTC)" and about[10] == []
    assert about[11:] == [["Column", "Type", "Description"], ["restaurant", "text", "The restaurant."],
                          ["orders", "whole number", "How many orders."], ["revenue", "number", "What they cost."]]
    assert leftovers(tmp_path) == []


def test_the_header_of_a_sheet_is_bold_and_frozen_and_has_a_filter(layer, tmp_path):
    serve.export_table(layer, "gold_revenue", tmp_path / "revenue.xlsx")

    sheet = openpyxl.load_workbook(tmp_path / "revenue.xlsx")["gold_revenue"]

    assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref == "A1:C3"
    assert [cell.font.bold for cell in sheet[1]] == [True, True, True]
    assert [cell.font.bold for cell in sheet[2]] == [False, False, False]
    assert [sheet.column_dimensions[letter].width for letter in "ABC"] == [12, 10, 10]  # a name and 2, but 10 at least
    labels = openpyxl.load_workbook(tmp_path / "revenue.xlsx")["about"]
    assert [labels[f"A{row}"].font.bold for row in (1, 5, 9)] == [True, True, True]
    assert [labels[f"A{row}"].font.bold for row in (12, 13)] == [True, False]  # the header of the columns, not a column


def test_every_kind_of_value_is_a_cell_that_a_person_can_read_and_a_text_is_never_a_formula(types, tmp_path):
    serve.export_table(types, "gold_types", tmp_path / "types.xlsx")

    sheet = openpyxl.load_workbook(tmp_path / "types.xlsx")["gold_types"]
    first, second = [[cell.value for cell in row] for row in sheet.iter_rows(min_row=2, max_row=3)]

    assert first[:6] == [1, "north", dt.datetime(2026, 10, 1), dt.datetime(2026, 10, 1, 8, 30),
                         dt.datetime(2026, 10, 1, 8, 30), 12.5]
    assert first[6:] == ["[1, 2]", '{"x": 1}', "=1+1", None, "123456789012345678"]
    assert second[1] == 'say "hi", ok' and second[3] == dt.datetime(2026, 10, 2, 7, 0)  # 09:00+02:00 is 07:00 UTC
    assert second[4] == dt.datetime(2026, 10, 2, 9, 0)  # a time with no zone is left as it is
    formula = sheet["I2"]
    assert (formula.value, formula.data_type) == ("=1+1", "s")
    described = {row[0].value: row[2] for row in openpyxl.load_workbook(tmp_path / "types.xlsx")["about"].iter_rows(
        min_row=13)}
    assert (described["formula"].value, described["formula"].data_type) == ("=1+1 is text, and not a sum.", "s")
    assert sheet["K2"].data_type == "s" and sheet["A2"].data_type == "n"
    assert sheet["C2"].number_format != "General" and sheet["D2"].number_format != "General"


def test_a_value_that_a_sheet_cannot_hold_is_changed_or_refused_and_never_corrupts_the_file(tmp_path):
    from openpyxl import Workbook

    sheet = Workbook(write_only=True).create_sheet("t")
    cell = serve._cell

    assert cell(sheet, float("nan"), "x") is None and cell(sheet, float("inf"), "x") is None
    assert cell(sheet, decimal.Decimal("NaN"), "x") is None and cell(sheet, decimal.Decimal("-Infinity"), "x") is None
    assert cell(sheet, decimal.Decimal("1.50"), "x") == decimal.Decimal("1.50")
    assert cell(sheet, 2 ** 53, "x") == 2 ** 53 and cell(sheet, 2 ** 53 + 1, "x").value == str(2 ** 53 + 1)
    assert cell(sheet, "bad \x01 char", "x").value == "bad  char"
    assert cell(sheet, b"\x00\xff", "x").value == "00ff"
    assert cell(sheet, dt.datetime(2026, 10, 1, 9, tzinfo=dt.timezone(dt.timedelta(hours=2))), "x") == dt.datetime(
        2026, 10, 1, 7)
    assert cell(sheet, dt.time(8, 30, tzinfo=UTC), "x") == dt.time(8, 30)
    assert cell(sheet, None, "x") is None and cell(sheet, True, "x") is True
    with pytest.raises(ConfigError, match=r"column 'x' has 32768 characters, and an Excel cell holds 32767"):
        cell(sheet, "x" * 32_768, "x")
    assert len(cell(sheet, "x" * 32_767, "x").value) == 32_767


def test_a_table_over_the_rows_of_a_sheet_is_refused_before_it_is_read_with_a_hint_for_csv(
        layer, tmp_path, monkeypatch):
    monkeypatch.setattr(serve, "XLSX_ROWS", 3)  # a sheet of 3 rows holds the header and 2 rows below it

    serve.export_table(layer, "gold_revenue", tmp_path / "fits.xlsx")
    monkeypatch.setattr(serve, "XLSX_ROWS", 2)
    with pytest.raises(ConfigError, match=r"gold\.gold_revenue has 2 rows, and an Excel sheet holds 1 below its "
                                          r"header\. Export CSV or Parquet instead"):
        serve.export_table(layer, "gold_revenue", tmp_path / "toolong.xlsx")

    assert (tmp_path / "fits.xlsx").exists() and not (tmp_path / "toolong.xlsx").exists()
    assert leftovers(tmp_path) == []


def test_a_table_over_the_rows_of_a_sheet_is_refused_before_a_row_is_read(layer, tmp_path, monkeypatch):
    from local_data_platform.format.iceberg import Iceberg

    def must_not_read(self, *args, **kwargs):
        raise AssertionError("the table was read")

    monkeypatch.setattr(Iceberg, "get", must_not_read)
    monkeypatch.setattr(serve, "XLSX_ROWS", 2)  # 2 rows and the header are 3

    with pytest.raises(ConfigError, match=r"has 2 rows, and an Excel sheet holds 1 below its header"):
        serve.export_table(layer, "gold_revenue", tmp_path / "toolong.xlsx")
    monkeypatch.setattr(serve, "XLSX_ROWS", 3)  # and a sheet of 3 rows holds them
    with pytest.raises(AssertionError, match="the table was read"):
        serve.export_table(layer, "gold_revenue", tmp_path / "fits.xlsx")


def test_the_rows_of_a_sheet_are_checked_again_when_they_cannot_be_counted_first(layer, tmp_path, monkeypatch):
    from pyiceberg.table import DataScan

    def cannot_count(self):
        raise RuntimeError("no count")

    monkeypatch.setattr(serve, "XLSX_ROWS", 2)
    monkeypatch.setattr(DataScan, "count", cannot_count)

    with pytest.raises(ConfigError, match=r"has 2 rows, and an Excel sheet holds 1 below its header"):
        serve.export_table(layer, "gold_revenue", tmp_path / "toolong.xlsx")

    assert not (tmp_path / "toolong.xlsx").exists() and leftovers(tmp_path) == []


@pytest.mark.filterwarnings("error::pytest.PytestUnraisableExceptionWarning")
def test_a_table_over_the_columns_and_a_value_over_the_cell_are_refused_and_write_nothing(
        layer, tmp_path, monkeypatch):
    monkeypatch.setattr(serve, "XLSX_COLUMNS", 3)
    serve.export_table(layer, "gold_revenue", tmp_path / "exactly.xlsx")  # a sheet of 3 columns holds 3
    monkeypatch.setattr(serve, "XLSX_COLUMNS", 2)
    with pytest.raises(ConfigError, match=r"has 3 columns, and an Excel sheet holds 2\. Export CSV or Parquet"):
        serve.export_table(layer, "gold_revenue", tmp_path / "wide.xlsx")
    monkeypatch.setattr(serve, "XLSX_COLUMNS", 100)
    monkeypatch.setattr(serve, "XLSX_CELL_CHARS", 4)
    with pytest.raises(ConfigError, match=r"a value in the column 'restaurant' has 5 characters, and an Excel cell "
                                          r"holds 4"):
        serve.export_table(layer, "gold_revenue", tmp_path / "long.xlsx")

    gc.collect()  # the sheet that was being written is finalised here, and must not write to a closed file
    assert (tmp_path / "exactly.xlsx").exists()
    assert not (tmp_path / "wide.xlsx").exists() and not (tmp_path / "long.xlsx").exists()
    assert leftovers(tmp_path) == []


def test_without_openpyxl_the_command_says_which_extra_to_install_before_it_reads_a_row(layer, tmp_path, monkeypatch):
    from local_data_platform.format.iceberg import Iceberg

    def must_not_read(self, *args, **kwargs):
        raise AssertionError("the table was read")

    monkeypatch.setitem(__import__("sys").modules, "openpyxl", None)
    monkeypatch.setattr(Iceberg, "get", must_not_read)

    with pytest.raises(EngineNotFound, match=r'pip install "local-data-platform\[excel\]"'):
        serve.export_table(layer, "gold_revenue", tmp_path / "revenue.xlsx")

    assert not (tmp_path / "revenue.xlsx").exists()


def test_an_excel_file_is_written_to_a_local_path_and_not_to_a_uri(layer):
    with pytest.raises(ConfigError, match=r"an XLSX file is written to a local path, and s3://bucket/x\.xlsx is not "
                                          r"one\. Export CSV or Parquet to a URI"):
        serve.export_table(layer, "gold_revenue", "s3://bucket/x.xlsx")


def test_the_sheet_of_a_table_has_a_name_that_excel_takes():
    assert serve._sheet_name("gold.daily_sales") == "daily_sales"
    assert serve._sheet_name("gold.about") == "about_data"
    assert serve._sheet_name("gold.ABOUT") == "ABOUT_data"
    assert serve._sheet_name("gold." + "x" * 40) == "x" * 31
    assert serve._sheet_name("gold.a[b]:c*d?e/f\\g") == "a_b__c_d_e_f_g"


def test_a_table_called_about_has_its_rows_on_a_sheet_of_another_name(tmp_path):
    silver(tmp_path)
    gold(tmp_path, "about")
    build_gold(tmp_path, release=True)

    serve.export_table(tmp_path, "about", tmp_path / "about.xlsx")

    assert openpyxl.load_workbook(tmp_path / "about.xlsx").sheetnames == ["about_data", "about"]


def test_the_excel_source_of_the_library_reads_the_file_back(layer, tmp_path):
    serve.export_table(layer, "gold_revenue", tmp_path / "revenue.xlsx")

    back = Excel("revenue", tmp_path / "revenue.xlsx", sheet="gold_revenue").get()

    assert back.to_pylist() == REVENUE


@pytest.mark.parametrize("kind, expected", [
    (pa.int32(), "whole number"), (pa.float64(), "number"), (pa.decimal128(10, 2), "decimal number"),
    (pa.bool_(), "true or false"), (pa.date32(), "date"), (pa.time64("us"), "time"), (pa.string(), "text"),
    (pa.large_string(), "text"), (pa.timestamp("us"), "date and time"), (pa.timestamp("us", tz="UTC"),
                                                                         "date and time (UTC)"),
    (pa.list_(pa.int64()), "a list or a record, written as JSON"), (pa.struct([("x", pa.int64())]),
                                                                    "a list or a record, written as JSON"),
    (pa.binary(), "binary")])
def test_a_column_is_described_in_words_for_a_person(kind, expected):
    assert serve._type_name(kind) == expected


# --------------------------------------------------------------------------- the DuckDB file


def test_the_duckdb_file_has_the_tables_of_the_release_with_their_comments_and_two_tables_of_its_own(layer, tmp_path):
    out = tmp_path / "gold.duckdb"

    done = serve.write_duckdb_file(layer, out)

    assert (done.path, done.release, done.tables, done.rows, done.metrics, done.left_out) == (
        out, "r1", ("gold_revenue",), 2, 2, ())
    assert query(out, "SELECT table_name FROM duckdb_tables() ORDER BY 1") == [
        ("_catalog",), ("_metrics",), ("gold_revenue",)]
    assert query(out, "SELECT * FROM gold_revenue ORDER BY 1") == [("north", 2, 30.5), ("south", 1, 7.25)]
    assert query(out, "SELECT comment FROM duckdb_tables() WHERE table_name = 'gold_revenue'") == [
        ("The gold_revenue table.",)]
    assert query(out, "SELECT column_name, comment FROM duckdb_columns() WHERE table_name = 'gold_revenue' "
                      "ORDER BY column_index") == [("restaurant", "The restaurant."), ("orders", "How many orders."),
                                                   ("revenue", "What they cost.")]
    assert leftovers(tmp_path) == []


def test_the_file_is_named_gold_and_still_has_no_schema_that_the_name_makes_ambiguous(layer, tmp_path):
    out = tmp_path / "gold.duckdb"
    serve.write_duckdb_file(layer, out)

    assert query(out, "SELECT count(*) FROM gold_revenue") == [(2,)]
    assert query(out, "SELECT count(*) FROM main.gold_revenue") == [(2,)]
    assert query(out, "SELECT count(*) FROM duckdb_schemas() WHERE schema_name = 'gold'") == [(0,)]
    assert query(out, "SELECT current_database()") == [("gold",)]


def test_the_catalog_table_has_a_row_for_each_table_with_its_release_snapshot_and_build_time(layer, tmp_path):
    out = tmp_path / "gold.duckdb"
    serve.write_duckdb_file(layer, out)
    released = releases.current_release(layer).table("gold.gold_revenue")
    snapshot = table(layer, "gold_revenue").table().snapshot_by_id(released.snapshot_id)

    columns = query(out, "DESCRIBE _catalog")
    rows = query(out, "SELECT * FROM _catalog")

    assert [(name, kind) for name, kind, *_ in columns] == [
        ("table_name", "VARCHAR"), ("description", "VARCHAR"), ("owner", "VARCHAR"), ("grain", "VARCHAR"),
        ("release", "VARCHAR"), ("snapshot_id", "BIGINT"), ("built_at", "TIMESTAMP")]
    assert rows == [("gold_revenue", "The gold_revenue table.", "finance", "one row per restaurant", "r1",
                     released.snapshot_id, dt.datetime.fromtimestamp(snapshot.timestamp_ms / 1000, UTC).replace(
                         tzinfo=None))]


def test_the_metrics_table_has_the_definitions_and_the_sql_that_computes_each(layer, tmp_path):
    out = tmp_path / "gold.duckdb"
    serve.write_duckdb_file(layer, out)

    rows = query(out, "SELECT metric, description, owner, type, format, table_name, definition, dimensions, "
                      "time_column, grains FROM _metrics ORDER BY metric")

    assert rows == [
        ("avg_ticket", "Revenue per order.", "kitchen", "ratio", "number", "gold_revenue",
         "sum(revenue) / NULLIF(sum(orders), 0)", "restaurant", None, ""),
        ("revenue", "What was sold.", "finance", "simple", "currency", "gold_revenue", "sum(revenue)", "restaurant",
         None, "")]
    assert query(out, "SELECT comment FROM duckdb_columns() WHERE table_name = '_metrics' AND column_name = "
                      "'definition'") == [("The SQL that computes it from that table, for the rows that are asked "
                                          "about.",)]


def test_a_folder_with_no_metrics_file_has_an_empty_metrics_table_with_the_same_columns(tmp_path):
    silver(tmp_path)
    gold(tmp_path)
    build_gold(tmp_path, release=True)
    out = tmp_path / "gold.duckdb"

    done = serve.write_duckdb_file(tmp_path, out)

    assert done.metrics == 0 and query(out, "SELECT count(*) FROM _metrics") == [(0,)]
    assert [row[0] for row in query(out, "DESCRIBE _metrics")] == list(serve._METRICS_COLUMNS)


def test_a_metrics_file_that_is_not_valid_stops_the_file_and_leaves_the_one_that_was_there(layer, tmp_path):
    out = tmp_path / "gold.duckdb"
    serve.write_duckdb_file(layer, out)
    before = out.read_bytes()
    write(layer / "metrics.json", {**METRICS, "metrics": [{**METRICS["metrics"][0], "format": "euro"}]})

    with pytest.raises(ConfigError, match=r"the metrics of .* are not valid, so they cannot go into the file"):
        serve.write_duckdb_file(layer, out)

    assert out.read_bytes() == before and leftovers(tmp_path) == []


def test_the_file_opens_with_no_extension_and_in_a_storage_version_that_duckdb_1_opens(layer, tmp_path):
    out = tmp_path / "gold.duckdb"
    serve.write_duckdb_file(layer, out)
    reference = tmp_path / "reference.duckdb"
    connection = duckdb.connect()
    try:
        connection.execute(f"ATTACH '{reference}' AS ref (STORAGE_VERSION 'v1.0.0')")
        connection.execute("CREATE TABLE ref.main.t AS SELECT 1 AS a")
        connection.execute("DETACH ref")
    except duckdb.Error:
        pytest.skip("this DuckDB writes one storage version, and cannot be asked for an older one")
    finally:
        connection.close()

    def storage_version(path: Path) -> int:
        header = path.read_bytes()[:32]
        assert header[8:12] == b"DUCK"
        return struct.unpack("<Q", header[12:20])[0]

    assert serve.STORAGE_VERSION == "v1.0.0" and storage_version(out) == storage_version(reference)
    settings = dict(query(out, "SELECT name, value FROM duckdb_settings() WHERE name LIKE 'auto%known_extensions'"))
    assert settings == {"autoinstall_known_extensions": "false", "autoload_known_extensions": "false"}
    assert query(out, "SELECT count(*) FROM duckdb_extensions() WHERE loaded AND extension_name = 'iceberg'") == [(0,)]


def test_a_failure_while_writing_leaves_the_file_that_was_there_and_no_temporary_file(layer, tmp_path, monkeypatch):
    out = tmp_path / "gold.duckdb"
    serve.write_duckdb_file(layer, out)
    before = out.read_bytes()
    more_orders(layer)
    build_gold(layer, release=True)
    created = []
    real = serve._create

    def fails_on_the_second(*args, **kwargs):
        created.append(args[1])
        if len(created) == 2:
            raise RuntimeError("the disk is full")
        return real(*args, **kwargs)

    monkeypatch.setattr(serve, "_create", fails_on_the_second)
    with pytest.raises(RuntimeError, match="disk is full"):
        serve.write_duckdb_file(layer, out)

    assert created == ["gold_revenue", "_catalog"]
    assert out.read_bytes() == before and leftovers(tmp_path) == []
    assert query(out, "SELECT count(*) FROM gold_revenue") == [(2,)]


def test_the_file_is_replaced_in_one_step_so_a_reader_has_the_old_file_or_the_new_one(layer, tmp_path, monkeypatch):
    out = tmp_path / "gold.duckdb"
    serve.write_duckdb_file(layer, out)
    more_orders(layer)
    build_gold(layer, release=True)
    seen = []
    real = os.replace

    def watching(source, target):
        seen.append((Path(source).name.startswith(".gold.duckdb.") and Path(source).name.endswith(".tmp"),
                     query(out, "SELECT count(*) FROM gold_revenue")))
        return real(source, target)

    monkeypatch.setattr(serve.os, "replace", watching)
    serve.write_duckdb_file(layer, out)

    assert seen == [(True, [(2,)])]  # at the moment of the rename the reader still has the old file
    assert query(out, "SELECT count(*) FROM gold_revenue") == [(3,)]


def test_a_file_that_a_program_holds_open_on_windows_is_an_error_that_says_how_to_write_a_new_name(
        layer, tmp_path, monkeypatch):
    out = tmp_path / "gold.duckdb"

    def denied(source, target):
        raise PermissionError("the file is in use")

    monkeypatch.setattr(serve.os, "replace", denied)
    with pytest.raises(ConfigError, match=r"cannot be replaced, because another program has it open, which Windows "
                                          r"does not allow\. Close it, or write a file with a new name for each "
                                          r"release, such as .*gold_\{release\}\.duckdb"):
        serve.write_duckdb_file(layer, out)

    assert not out.exists() and leftovers(tmp_path) == []


def test_an_excel_file_that_a_program_holds_open_is_refused_with_the_same_hint_and_leaves_no_temporary_file(
        layer, tmp_path, monkeypatch):
    def denied(source, target):
        raise PermissionError("the file is in use")

    monkeypatch.setattr(serve.os, "replace", denied)

    with pytest.raises(ConfigError, match=r"revenue\.xlsx cannot be replaced, because another program has it open"):
        serve.export_table(layer, "gold_revenue", tmp_path / "revenue.xlsx")

    assert not (tmp_path / "revenue.xlsx").exists() and leftovers(tmp_path) == []


def test_release_in_the_name_of_the_file_is_the_release_and_live_when_there_is_none(layer, tmp_path):
    first = serve.write_duckdb_file(layer, tmp_path / "gold_{release}.duckdb")
    live = serve.write_duckdb_file(layer, tmp_path / "gold_{release}.duckdb", live=True)

    assert (first.path.name, live.path.name) == ("gold_r1.duckdb", "gold_live.duckdb")
    assert serve.expand("a/{release}/x.duckdb", "r3") == Path("a/r3/x.duckdb")
    assert serve.expand("x.duckdb", None) == Path("x.duckdb")
    assert query(live.path, "SELECT release FROM _catalog") == [(None,)] and live.release is None
    current = table(layer, "gold_revenue").table().current_snapshot()
    assert query(live.path, "SELECT snapshot_id, built_at IS NOT NULL FROM _catalog") == [(current.snapshot_id, True)]


def test_the_file_holds_the_release_that_is_asked_for_and_the_tables_as_they_are_now_with_live(layer, tmp_path):
    more_orders(layer)
    build_gold(layer, release=True)  # r2

    older = serve.write_duckdb_file(layer, tmp_path / "r1.duckdb", release="r1")
    newer = serve.write_duckdb_file(layer, tmp_path / "r2.duckdb")

    assert (older.release, older.rows, newer.release, newer.rows) == ("r1", 2, "r2", 3)
    assert query(tmp_path / "r1.duckdb", "SELECT DISTINCT release FROM _catalog") == [("r1",)]
    assert query(tmp_path / "r2.duckdb", "SELECT DISTINCT release FROM _catalog") == [("r2",)]
    assert older.source.startswith("release r1, cut ")


def test_a_gold_table_that_the_release_does_not_hold_is_left_out_and_named(layer, tmp_path):
    run_config(gold(layer, "gold_newer"))

    done = serve.write_duckdb_file(layer, tmp_path / "gold.duckdb")

    assert done.tables == ("gold_revenue",) and done.left_out == ("gold.gold_newer",)
    assert "Left out, because r1 does not hold them, gold.gold_newer" in done.summary()
    assert query(tmp_path / "gold.duckdb", "SELECT table_name FROM _catalog") == [("gold_revenue",)]
    live = serve.write_duckdb_file(layer, tmp_path / "live.duckdb", live=True)
    assert live.tables == ("gold_newer", "gold_revenue") and live.left_out == ()


@pytest.mark.parametrize("name", ["_catalog", "_metrics"])
def test_a_gold_table_called_like_a_table_of_the_file_is_refused_because_the_file_uses_the_name(tmp_path, name):
    silver(tmp_path)
    gold(tmp_path, name)
    build_gold(tmp_path, release=True)

    with pytest.raises(ConfigError, match=rf"a gold table is called {name}, which the DuckDB file uses for its own "
                                          r"table"):
        serve.write_duckdb_file(tmp_path, tmp_path / "gold.duckdb")

    assert not (tmp_path / "gold.duckdb").exists() and leftovers(tmp_path) == []


def test_every_kind_of_value_is_in_the_file_as_the_type_that_the_table_has(types, tmp_path):
    out = tmp_path / "gold.duckdb"
    serve.write_duckdb_file(types, out)

    described = {name: kind for name, kind, *_ in query(out, "DESCRIBE gold_types")}
    connection = open_file(out)
    try:
        connection.execute("SET TimeZone = 'UTC'")
        rows = connection.execute(
            "SELECT id, name, CAST(day AS VARCHAR), CAST(moment AS VARCHAR), CAST(local_moment AS VARCHAR), "
            "CAST(amount AS VARCHAR), CAST(tags AS VARCHAR), CAST(info AS VARCHAR), formula, ratio, big "
            "FROM gold_types ORDER BY id").fetchall()
    finally:
        connection.close()

    assert described["day"] == "DATE" and described["amount"] == "DECIMAL(10,2)" and described["big"] == "BIGINT"
    assert described["moment"] == "TIMESTAMP WITH TIME ZONE" and described["local_moment"] == "TIMESTAMP"
    assert described["tags"] == "INTEGER[]" and described["info"].startswith("STRUCT(")
    assert rows[0] == (1, "north", "2026-10-01", "2026-10-01 08:30:00+00", "2026-10-01 08:30:00", "12.50", "[1, 2]",
                       "{'x': 1}", "=1+1", None, 123456789012345678)
    assert rows[1][1] == 'say "hi", ok' and rows[1][3] == "2026-10-02 07:00:00+00"
    assert rows[1][4] == "2026-10-02 09:00:00"


def test_a_file_that_cannot_be_written_is_still_read_for_its_release(layer, tmp_path):
    out = tmp_path / "gold.duckdb"
    serve.write_duckdb_file(layer, out)
    out.chmod(0o444)
    try:
        assert serve.file_release(out) == "r1"
    finally:
        out.chmod(0o644)


def test_the_release_that_a_file_holds_is_read_from_its_catalog_and_a_bad_file_holds_none(layer, tmp_path):
    out = tmp_path / "gold.duckdb"
    assert serve.file_release(out) is None
    serve.write_duckdb_file(layer, out)
    assert serve.file_release(out) == "r1"
    serve.write_duckdb_file(layer, tmp_path / "live.duckdb", live=True)
    assert serve.file_release(tmp_path / "live.duckdb") is None
    junk = tmp_path / "junk.duckdb"
    junk.write_bytes(b"not a database")
    assert serve.file_release(junk) is None
    plain = tmp_path / "plain.duckdb"
    connection = duckdb.connect(str(plain))
    connection.execute("CREATE TABLE t AS SELECT 1 AS a")
    connection.execute("CREATE TABLE _catalog AS SELECT * FROM (VALUES ('r1'), ('r2')) AS t(release)")
    connection.close()
    assert serve.file_release(plain) is None  # a catalog of two releases is not the file of one


def test_a_duckdb_that_has_no_storage_version_option_writes_its_own_version(tmp_path):
    class Connection:
        def __init__(self):
            self.statements = []

        def execute(self, sql):
            self.statements.append(sql)
            if "STORAGE_VERSION" in sql:
                raise duckdb.Error("Unrecognized option for attach storage_version")

    connection = Connection()
    serve._attach(connection, duckdb, tmp_path / "x.duckdb")

    assert connection.statements == [
        f"ATTACH '{tmp_path / 'x.duckdb'}' AS ldp_out (STORAGE_VERSION 'v1.0.0')",
        f"ATTACH '{tmp_path / 'x.duckdb'}' AS ldp_out"]


def test_a_comment_on_a_column_that_the_table_does_not_have_is_left_out_and_is_no_error():
    connection = duckdb.connect()
    connection.execute("ATTACH ':memory:' AS ldp_out")
    try:
        serve._create(connection, "t", pa.table({"a": [1]}), "A table.", {"a": "A column.", "gone": "Not a column."})

        assert connection.execute("SELECT column_name, comment FROM duckdb_columns() WHERE table_name = 't'"
                                  ).fetchall() == [("a", "A column.")]
    finally:
        connection.close()


def test_what_a_failed_write_leaves_behind_is_the_file_and_its_log_and_nothing_else(tmp_path):
    (tmp_path / ".x.duckdb.1.tmp").write_bytes(b"half")
    (tmp_path / ".x.duckdb.1.tmp.wal").write_bytes(b"log")
    (tmp_path / "x.duckdb").write_bytes(b"the file that was there")

    serve._remove(tmp_path / ".x.duckdb.1.tmp")
    serve._remove(tmp_path / ".x.duckdb.2.tmp")  # a file that is not there is no error

    assert sorted(item.name for item in tmp_path.iterdir()) == ["x.duckdb"]


def test_without_duckdb_the_command_says_which_extra_to_install(layer, tmp_path, monkeypatch):
    monkeypatch.setitem(__import__("sys").modules, "duckdb", None)

    with pytest.raises(EngineNotFound, match=r'pip install "local-data-platform\[duckdb\]"'):
        serve.write_duckdb_file(layer, tmp_path / "gold.duckdb")

    assert not (tmp_path / "gold.duckdb").exists()


def test_text_in_any_script_and_a_table_with_no_rows_go_through_every_format(tmp_path):
    silver(tmp_path)
    gold(tmp_path, "gold_words", sql="SELECT 'caf\u00e9 \u6771\u4eac \U0001F355' AS word, 1 AS n",
         columns={"word": "A word.", "n": "A number."})
    emptied = table(tmp_path, "gold_none")  # a build writes no empty table, so it is written and then emptied
    emptied.put(pa.table({"word": ["x"], "n": pa.array([1], pa.int64())}))
    emptied.table().delete("n >= 0")
    metadata = {"source": {"name": "gold_none", "format": "ICEBERG", "catalog": GOLD_CATALOG},
                "target": {"name": "gold_none", "format": "ICEBERG", "catalog": GOLD_CATALOG},
                "docs": {"layer": "gold", "description": "Nothing yet.", "grain": "one row per word",
                         "columns": {"word": "A word.", "n": "A number."}}}
    write(tmp_path / "gold_none.json", {"identifier": "gold_none", "who": "ops", "metadata": metadata})
    assert build_gold(tmp_path, release=True).exit_code == 0
    out = tmp_path / "out"

    for name in ("csv", "parquet", "jsonl", "xlsx"):
        serve.export_table(tmp_path, "gold_words", out / f"words.{name}")
        serve.export_table(tmp_path, "gold_none", out / f"none.{name}")
    serve.write_duckdb_file(tmp_path, out / "gold.duckdb")

    word = "caf\u00e9 \u6771\u4eac \U0001F355"
    assert csv_rows(out / "words.csv") == [{"word": word, "n": 1}]
    assert pq.read_table(out / "words.parquet").to_pylist() == [{"word": word, "n": 1}]
    assert json.loads((out / "words.jsonl").read_text()) == {"word": word, "n": 1}
    book = openpyxl.load_workbook(out / "words.xlsx")
    assert [[cell.value for cell in row] for row in book["gold_words"].iter_rows()] == [["word", "n"], [word, 1]]
    assert query(out / "gold.duckdb", "SELECT word, n FROM gold_words") == [(word, 1)]

    assert pacsv.read_csv(out / "none.csv").column_names == ["word", "n"] and csv_rows(out / "none.csv") == []
    assert pq.read_table(out / "none.parquet").num_rows == 0 and (out / "none.jsonl").read_text() == ""
    empty = openpyxl.load_workbook(out / "none.xlsx")
    assert [[cell.value for cell in row] for row in empty["gold_none"].iter_rows()] == [["word", "n"]]
    assert empty["about"]["B9"].value == 0  # the rows of the table, which is the ninth line of the sheet
    assert query(out / "gold.duckdb", "SELECT count(*) FROM gold_none") == [(0,)]
    assert [name for name, *_ in query(out / "gold.duckdb", "DESCRIBE gold_none")] == ["word", "n"]


def test_a_comment_and_a_name_with_quotes_are_written_as_they_are(tmp_path):
    silver(tmp_path)
    columns = {"a b": "It's \"quoted\".", "it's": "A 'name'.", 'say "hi"': "A name with double quotes."}
    gold(tmp_path, "gold_quotes", sql="SELECT 1 AS \"a b\", 2 AS \"it's\", 3 AS \"say \"\"hi\"\"\"", columns=columns,
         docs={"description": "Don't break; DROP TABLE x --"})
    build_gold(tmp_path, release=True)

    serve.write_duckdb_file(tmp_path, tmp_path / "gold.duckdb")

    assert query(tmp_path / "gold.duckdb", "SELECT comment FROM duckdb_tables() WHERE table_name = 'gold_quotes'") == [
        ("Don't break; DROP TABLE x --",)]
    assert query(tmp_path / "gold.duckdb", "SELECT column_name, comment FROM duckdb_columns() WHERE table_name = "
                                           "'gold_quotes' ORDER BY column_index") == [
        ("a b", "It's \"quoted\"."), ("it's", "A 'name'."), ('say "hi"', "A name with double quotes.")]


# --------------------------------------------------------------------------- ldp gold duckdb


def test_ldp_gold_duckdb_writes_the_file_and_says_what_it_holds(layer, tmp_path, capsys):
    out = tmp_path / "gold.duckdb"

    code, text, err = cli(capsys, "gold", "duckdb", layer, "--out", out)

    assert code == 0 and err == ""
    assert text.startswith(f"wrote {out}, 1 table and 2 metrics, from release r1, cut ") and text.endswith(" KB\n")
    assert query(out, "SELECT count(*) FROM gold_revenue") == [(2,)]


def test_ldp_gold_duckdb_takes_a_release_or_live_and_prints_json(layer, tmp_path, capsys):
    more_orders(layer)
    build_gold(layer, release=False)

    code, text, _ = cli(capsys, "gold", "duckdb", layer, "--out", tmp_path / "live.duckdb", "--live", "--json")
    data = json.loads(text)
    assert code == 0 and (data["release"], data["tables"], data["rows"], data["metrics"], data["left_out"]) == (
        None, ["gold_revenue"], 3, 2, [])
    assert data["path"] == str(tmp_path / "live.duckdb") and data["size"] > 0
    code, _, _ = cli(capsys, "gold", "duckdb", layer, "--out", tmp_path / "r1.duckdb", "--release", "r1")
    assert code == 0 and query(tmp_path / "r1.duckdb", "SELECT count(*) FROM gold_revenue") == [(2,)]
    code, _, err = cli(capsys, "gold", "duckdb", layer, "--out", tmp_path / "x.duckdb", "--release", "r1", "--live")
    assert code == 1 and "not allowed with argument" in err and not (tmp_path / "x.duckdb").exists()


def test_ldp_gold_duckdb_needs_an_out_and_a_folder_with_a_gold_table(layer, tmp_path, capsys):
    code, _, err = cli(capsys, "gold", "duckdb", layer)
    assert code == 1 and "--out" in err

    empty = tmp_path / "empty"
    empty.mkdir()
    code, _, err = cli(capsys, "gold", "duckdb", empty, "--out", tmp_path / "x.duckdb")
    assert code == 1 and "has no gold config" in err and not (tmp_path / "x.duckdb").exists()


# --------------------------------------------------------------------------- ldp gold build --duckdb


def test_a_build_with_duckdb_writes_the_file_after_it_cuts_the_release(tmp_path, capsys):
    silver(tmp_path)
    gold(tmp_path)
    out = tmp_path / "out" / "gold.duckdb"

    code, text, err = cli(capsys, "gold", "build", tmp_path, "--duckdb", out)

    assert code == 0 and err == ""
    assert "release r1 cut with 1 table and 2 rows in gold" in text
    line = next(item for item in text.splitlines() if item.startswith("duckdb: "))
    assert line.startswith(f"duckdb: wrote {out}, 1 table, from release r1, cut ")
    assert serve.file_release(out) == "r1" and query(out, "SELECT count(*) FROM gold_revenue") == [(2,)]


def test_a_build_that_cuts_the_next_release_refreshes_the_file(layer, tmp_path, capsys):
    out = tmp_path / "gold.duckdb"
    serve.write_duckdb_file(layer, out)
    more_orders(layer)

    code, text, _ = cli(capsys, "gold", "build", layer, "--duckdb", out)

    assert code == 0 and "release r2 cut" in text and "duckdb: wrote " in text
    assert serve.file_release(out) == "r2" and query(out, "SELECT count(*) FROM gold_revenue") == [(3,)]


def test_a_build_that_cuts_nothing_leaves_a_file_of_the_current_release_and_writes_a_missing_one(
        checked, tmp_path, capsys):
    out = tmp_path / "out" / "gold.duckdb"
    code, text, _ = cli(capsys, "gold", "build", checked, "--duckdb", out)
    assert code == 0 and "release r1 cut" in text and "duckdb: wrote " in text and serve.file_release(out) == "r1"

    before = out.read_bytes()
    code, text, _ = cli(capsys, "gold", "build", checked, "--duckdb", out)
    assert code == 0 and "no release cut, because the tables are at the snapshots of r1" in text
    assert f"duckdb: {out} was left as it is, because it already holds release r1" in text
    assert out.read_bytes() == before

    out.unlink()  # a file that is not there is written, though no release was cut
    code, text, _ = cli(capsys, "gold", "build", checked, "--duckdb", out)
    assert code == 0 and "no release cut" in text and "duckdb: wrote " in text and serve.file_release(out) == "r1"

    serve.write_duckdb_file(checked, out, live=True)  # a file that holds no release is not the file of the current one
    code, text, _ = cli(capsys, "gold", "build", checked, "--duckdb", out)
    assert code == 0 and "duckdb: wrote " in text and serve.file_release(out) == "r1"


def test_the_file_is_written_for_the_release_that_the_build_cut_and_not_for_one_that_was_cut_after(layer, tmp_path):
    out = tmp_path / "gold.duckdb"
    first = build_gold(layer, release=True)  # r2, because the fixture cut r1
    more_orders(layer)
    assert build_gold(layer, release=True).release.release.name == "r3"

    outcome = serve.refresh_after_build(layer, out, first)

    assert outcome.status == "written" and outcome.file.release == "r2" and serve.file_release(out) == "r2"
    assert query(out, "SELECT count(*) FROM gold_revenue") == [(2,)]


def test_a_release_just_cut_is_written_even_when_a_file_says_it_holds_that_release(
        layer, tmp_path, capsys, monkeypatch):
    out = tmp_path / "gold.duckdb"
    serve.write_duckdb_file(layer, out)
    more_orders(layer)
    monkeypatch.setattr(serve, "file_release", lambda path: "r2")  # a file of another lake, with a release of that name

    code, text, _ = cli(capsys, "gold", "build", layer, "--duckdb", out)

    assert code == 0 and "release r2 cut" in text and f"duckdb: wrote {out}" in text
    assert query(out, "SELECT count(*) FROM gold_revenue") == [(3,)]


def test_a_build_with_release_in_the_name_writes_a_file_for_each_release(layer, tmp_path, capsys):
    pattern = tmp_path / "gold_{release}.duckdb"
    cli(capsys, "gold", "build", layer, "--duckdb", pattern)  # r1 was cut by the fixture, so this cuts r2
    more_orders(layer)

    code, text, _ = cli(capsys, "gold", "build", layer, "--duckdb", pattern)

    assert code == 0
    assert sorted(item.name for item in tmp_path.glob("gold_*.duckdb")) == ["gold_r2.duckdb", "gold_r3.duckdb"]
    assert serve.file_release(tmp_path / "gold_r2.duckdb") == "r2"
    assert f"wrote {tmp_path / 'gold_r3.duckdb'}" in text


def test_a_build_that_fails_leaves_the_file_of_the_last_release_in_place(layer, tmp_path, capsys):
    out = tmp_path / "gold.duckdb"
    serve.write_duckdb_file(layer, out)
    before = out.read_bytes()
    gold(layer, "gold_revenue", sql="SELECT * FROM a_table_that_is_not_there")

    code, text, err = cli(capsys, "gold", "build", layer, "--duckdb", out)

    assert code == 1 and "no release cut" in text
    assert f"duckdb: {out} was left as it is, because the build did not end well" in text
    assert out.read_bytes() == before and serve.file_release(out) == "r1" and leftovers(tmp_path) == []
    assert releases.current_release(layer).name == "r1" and "gold_revenue" in err


def test_a_build_that_a_gate_blocks_leaves_the_file_in_place_and_exits_with_3(layer, tmp_path, capsys):
    out = tmp_path / "gold.duckdb"
    serve.write_duckdb_file(layer, out)
    before = out.read_bytes()
    more_orders(layer)
    gold(layer, "gold_revenue", quality={"checks": [{"check": "row_count", "min": 100}]})

    code, text, _ = cli(capsys, "gold", "build", layer, "--duckdb", out)

    assert code == 3 and "duckdb: " in text and "was left as it is, because the build did not end well" in text
    assert out.read_bytes() == before


def test_a_file_that_cannot_be_written_after_the_release_is_cut_is_an_error_and_the_old_file_stays(
        layer, tmp_path, capsys, monkeypatch):
    out = tmp_path / "gold.duckdb"
    serve.write_duckdb_file(layer, out)
    before = out.read_bytes()
    more_orders(layer)

    def full(*args, **kwargs):
        raise OSError("the disk is full")

    monkeypatch.setattr(serve, "write_duckdb_file", full)
    code, text, err = cli(capsys, "gold", "build", layer, "--duckdb", out)

    assert code == 1 and "release r2 cut" in text and f"duckdb: {out} was not written" in text
    assert f"ldp: the DuckDB file was not written, and {out} is as it was" in err and "the disk is full" in err
    assert out.read_bytes() == before and releases.current_release(layer).name == "r2"


def test_duckdb_cannot_be_used_with_no_release_and_nothing_is_built(tmp_path, capsys):
    silver(tmp_path)
    gold(tmp_path)

    code, text, err = cli(capsys, "gold", "build", tmp_path, "--no-release", "--duckdb", tmp_path / "gold.duckdb")

    assert code == 1 and text == ""
    assert "--duckdb writes the file for a release, so it cannot be used with --no-release" in err
    assert not list((tmp_path / "warehouse").glob("gold*")) and not (tmp_path / "gold.duckdb").exists()


def test_the_json_of_a_build_has_what_was_done_about_the_file(checked, tmp_path, capsys):
    out = tmp_path / "gold.duckdb"

    code, text, _ = cli(capsys, "gold", "build", checked, "--duckdb", out, "--json")

    data = json.loads(text)["duckdb"]
    assert code == 0 and (data["status"], data["path"], data["error"]) == ("written", str(out), None)
    assert data["file"]["release"] == "r1" and data["file"]["tables"] == ["elsewhere"]
    code, text, _ = cli(capsys, "gold", "build", checked, "--duckdb", out, "--json")
    kept = json.loads(text)["duckdb"]
    assert (kept["status"], kept["reason"], kept["file"]) == ("kept", "it already holds release r1", None)


def test_the_outcome_of_a_refresh_says_what_it_did_in_one_line(tmp_path):
    left = serve.FileOutcome(tmp_path / "g.duckdb", "left", "there is no release")
    failed = serve.FileOutcome(tmp_path / "g.duckdb", "failed", error="OSError: x")

    assert left.summary() == f"{tmp_path / 'g.duckdb'} was left as it is, because there is no release"
    assert failed.summary() == f"{tmp_path / 'g.duckdb'} was not written"
    assert left.to_dict() == {"path": str(tmp_path / "g.duckdb"), "status": "left", "reason": "there is no release",
                              "file": None, "error": None}


def test_the_datasets_export_and_the_gold_export_share_the_csv_writer(layer, tmp_path):
    assert "csv" in datasets.EXPORT_FORMATS and set(serve.EXPORT_FORMATS) == {"csv", "parquet", "jsonl", "xlsx"}
    serve.export_table(layer, "gold_revenue", tmp_path / "gold.csv")
    datasets.write_rows(table(layer, "gold_revenue").get(), tmp_path / "datasets.csv", "csv")

    assert (tmp_path / "gold.csv").read_bytes() == (tmp_path / "datasets.csv").read_bytes()
