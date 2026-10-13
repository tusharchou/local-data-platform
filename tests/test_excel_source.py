"""The XLSX source: a sheet of an Excel workbook read into Arrow, and the ExcelToCSV and ExcelToIceberg routes.

Every workbook is built with openpyxl when the test runs, so the tests need no fixture files and no network.
"""

import datetime as dt
import json
import logging
import re
import sys
import zipfile
from pathlib import Path

import openpyxl
import pyarrow as pa
import pytest
from openpyxl.styles import Font

from local_data_platform import Config, SupportedFormat
from local_data_platform.cli import main
from local_data_platform.events import dataset_ref
from local_data_platform.exceptions import ConfigError, DataQualityError, EngineNotFound
from local_data_platform.format import excel as excel_module
from local_data_platform.format.excel import INSTALL_HINT, Excel
from local_data_platform.format.iceberg import Iceberg
from local_data_platform.pipeline.builders import excel_from_config
from local_data_platform.pipeline.registry import create_pipeline, get_pipeline_class, make_route
from local_data_platform.spec import FORMATS, json_schema, plan, validate_spec

DAY = dt.date(2026, 7, 1)


def workbook(path: Path, sheets: dict[str, list[list]]) -> Path:
    """Write an ``.xlsx`` with one sheet for each item of ``sheets``: its title and its rows."""
    book = openpyxl.Workbook()
    book.remove(book.active)
    for title, rows in sheets.items():
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(row)
    book.save(path)
    return path


def with_cached_values(path: Path, values: dict[str, object]) -> None:
    """Give formula cells of the first sheet the value Excel would have saved with them (openpyxl saves none)."""
    temporary = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                text = data.decode()
                for cell, value in values.items():
                    text = re.sub(rf'(<c r="{cell}"[^>]*><f>[^<]*</f>)<v\s*/>', rf"\g<1><v>{value}</v>", text)
                data = text.encode()
            target.writestr(item, data)
    temporary.replace(path)


def read(path: Path, **options) -> pa.Table:
    return Excel("sheet", path, **options).get()


@pytest.fixture
def sales(tmp_path) -> Path:
    """A workbook with a title row, then a header row, and a second sheet."""
    return workbook(tmp_path / "sales.xlsx", {
        "Sales": [["Sales for July"], ["day", "dish", "orders", "revenue"],
                  [DAY, "pad thai", 12, 143.5], [dt.date(2026, 7, 2), "green curry", 9, 117]],
        "Menu": [["dish", "price"], ["pad thai", 11.5], ["green curry", 13]],
    })


# --------------------------------------------------------------------------- types


def test_each_column_gets_the_arrow_type_that_fits_its_cells(tmp_path):
    path = workbook(tmp_path / "types.xlsx", {"T": [
        ["whole", "number", "text", "flag", "day", "moment", "clock", "empty"],
        [1, 1.5, "a", True, DAY, dt.datetime(2026, 7, 1, 12, 30, 15), dt.time(9, 30), None],
        [2, 2, "b", False, dt.date(2026, 7, 2), dt.datetime(2026, 7, 2, 8, 0), dt.time(17, 45, 5), None],
        [None, None, None, None, None, None, None, None],
        [3, 4.25, "c", True, None, None, None, None],
    ]})

    table = read(path)

    assert dict(zip(table.schema.names, (str(field.type) for field in table.schema))) == {
        "whole": "int64", "number": "double", "text": "string", "flag": "bool", "day": "timestamp[us]",
        "moment": "timestamp[us]", "clock": "time64[us]", "empty": "string"}
    assert table.column("whole").to_pylist() == [1, 2, None, 3]
    assert table.column("number").to_pylist() == [1.5, 2.0, None, 4.25]
    assert table.column("flag").to_pylist() == [True, False, None, True]
    assert table.column("day").to_pylist() == [dt.datetime(2026, 7, 1), dt.datetime(2026, 7, 2), None, None]
    assert table.column("moment").to_pylist()[0] == dt.datetime(2026, 7, 1, 12, 30, 15)
    assert table.column("clock").to_pylist()[1] == dt.time(17, 45, 5)
    assert table.column("empty").to_pylist() == [None, None, None, None]  # nothing in it, so text of nulls


@pytest.mark.parametrize("values, expected_type, expected", [
    ([1, 2, 3], "int64", [1, 2, 3]),
    ([1.0, 2.0], "int64", [1, 2]),  # a float that is a whole number is a whole number
    ([1, 2.5], "double", [1.0, 2.5]),
    ([2**70], "double", [float(2**70)]),  # too large for int64
    ([-3, 4], "int64", [-3, 4]),
    ([dt.date(2026, 7, 1), dt.datetime(2026, 7, 2, 3, 4, 5)], "timestamp[us]",
     [dt.datetime(2026, 7, 1), dt.datetime(2026, 7, 2, 3, 4, 5)]),
])
def test_number_and_date_columns(tmp_path, values, expected_type, expected):
    path = workbook(tmp_path / "t.xlsx", {"T": [["v"], *[[value] for value in values]]})

    column = read(path).column("v")

    assert str(column.type) == expected_type
    assert column.to_pylist() == (pytest.approx(expected) if expected_type == "double" else expected)


@pytest.mark.parametrize("values, kinds, expected", [
    ([1, "n/a", 2.5], "number and text", ["1", "n/a", "2.5"]),
    ([True, 3], "bool and number", ["TRUE", "3"]),
    ([DAY, "soon"], "date and text", ["2026-07-01 00:00:00", "soon"]),  # Excel keeps a date as a point in time
    ([dt.datetime(2026, 7, 1, 8, 30), "late"], "date and text", ["2026-07-01 08:30:00", "late"]),
    ([dt.time(9, 5), 1], "number and time", ["1", "09:05:00"][::-1]),
])
def test_a_column_that_mixes_kinds_becomes_text_and_a_warning_names_it(tmp_path, caplog, values, kinds, expected):
    path = workbook(tmp_path / "mixed.xlsx", {"Data": [["id", "odd one"], *[[i, v] for i, v in enumerate(values)]]})
    caplog.set_level(logging.WARNING, logger="local_data_platform.format.excel")

    table = read(path)

    assert table.schema.field("id").type == pa.int64()  # the other columns keep their types
    assert table.schema.field("odd one").type == pa.string()
    assert table.column("odd one").to_pylist() == expected
    [warning] = [record.getMessage() for record in caplog.records]
    assert "column 'odd one' holds " in warning and kinds in warning and "(sheet 'Data')" in warning


def test_empty_cells_are_nulls_whatever_the_type(tmp_path):
    path = workbook(tmp_path / "t.xlsx", {"T": [["a", "b"], [1, None], [None, "x"], [3, None]]})

    table = read(path)

    assert table.column("a").to_pylist() == [1, None, 3] and table.column("b").to_pylist() == [None, "x", None]


# --------------------------------------------------------------------------- formulas


def test_a_formula_cell_gives_the_value_that_was_saved_with_it(tmp_path):
    path = workbook(tmp_path / "f.xlsx", {"T": [["a", "b", "total"], [3, 4, "=A2+B2"], [5, 6, "=A3+B3"]]})
    with_cached_values(path, {"C2": 7, "C3": 11})

    table = read(path)

    assert table.column("total").to_pylist() == [7, 11]  # not the text of the formula


def test_a_formula_that_was_never_calculated_is_a_null(tmp_path):
    path = workbook(tmp_path / "f.xlsx", {"T": [["a", "b", "total"], [3, 4, "=A2+B2"]]})  # openpyxl saves no value

    assert read(path).column("total").to_pylist() == [None]


# --------------------------------------------------------------------------- the sheet and the header row


def test_the_first_sheet_is_read_by_default(sales):
    table = read(sales, header_row=2)

    assert table.column_names == ["day", "dish", "orders", "revenue"] and table.num_rows == 2


@pytest.mark.parametrize("sheet", ["Menu", "menu", " MENU ", 1])
def test_a_sheet_is_found_by_name_without_regard_to_case_or_by_position(sales, sheet):
    table = read(sales, sheet=sheet)

    assert table.column_names == ["dish", "price"] and table.column("price").to_pylist() == [11.5, 13.0]


def test_the_rows_above_the_header_row_are_skipped(sales):
    table = read(sales, sheet="Sales", header_row=2)

    assert table.column_names == ["day", "dish", "orders", "revenue"]
    assert table.column("revenue").to_pylist() == [143.5, 117.0]
    with pytest.raises(ConfigError, match="the header cell B1 is empty"):  # the title row is not a header
        read(sales, sheet="Sales", header_row=1)


def test_a_sheet_that_does_not_exist_names_the_file_and_the_sheets_there_are(sales):
    with pytest.raises(ConfigError) as error:
        read(sales, sheet="Salez")

    assert str(error.value) == f"{sales}: there is no sheet named 'Salez'; the sheets are ['Sales', 'Menu']"


def test_a_position_that_is_past_the_last_sheet_names_the_file(sales):
    with pytest.raises(ConfigError) as error:
        read(sales, sheet=5)

    assert str(error.value) == f"{sales}: there is no sheet at position 5; the workbook has 2: ['Sales', 'Menu']"


def test_a_header_row_past_the_end_of_the_sheet_is_an_error(sales):
    with pytest.raises(ConfigError, match=r"\(sheet 'Menu'\): the sheet has no row 9, so there is no header to read"):
        read(sales, sheet="Menu", header_row=9)


def test_a_sheet_with_nothing_in_it_is_an_error(tmp_path):
    path = workbook(tmp_path / "empty.xlsx", {"Empty": []})

    with pytest.raises(ConfigError, match=r"\(sheet 'Empty'\): the sheet has no row 1"):
        read(path)


def test_a_header_row_with_nothing_in_it_is_an_error(tmp_path):
    path = workbook(tmp_path / "t.xlsx", {"T": [["a"], [None, None]]})

    with pytest.raises(ConfigError, match=r"row 2, the header row, is empty"):
        read(path, header_row=2)


# --------------------------------------------------------------------------- trailing empty rows and columns


def test_empty_rows_at_the_end_and_empty_columns_at_the_right_edge_are_dropped(tmp_path):
    path = workbook(tmp_path / "t.xlsx", {"T": [["a", "b"], [1, "x"], [2, "y"], [], [None, None], []]})
    book = openpyxl.load_workbook(path)
    book["T"]["H40"].font = Font(bold=True)  # formatting leaves empty cells far from the data
    book.save(path)

    table = read(path)

    assert table.column_names == ["a", "b"] and table.num_rows == 2


def test_an_empty_row_between_data_rows_is_kept_as_a_row_of_nulls(tmp_path):
    path = workbook(tmp_path / "t.xlsx", {"T": [["a", "b"], [1, "x"], [], [3, "z"]]})

    table = read(path)

    assert table.column("a").to_pylist() == [1, None, 3] and table.column("b").to_pylist() == ["x", None, "z"]


def test_a_column_with_a_header_and_no_data_is_kept(tmp_path):
    path = workbook(tmp_path / "t.xlsx", {"T": [["a", "b", "notes"], [1, "x", None]]})

    table = read(path)

    assert table.column_names == ["a", "b", "notes"] and table.column("notes").type == pa.string()


def test_a_sheet_with_a_header_and_no_rows_gives_an_empty_table(tmp_path):
    path = workbook(tmp_path / "t.xlsx", {"T": [["a", "b"]]})

    table = read(path)

    assert table.column_names == ["a", "b"] and table.num_rows == 0


def test_a_sheet_whose_recorded_size_is_wrong_is_still_read_whole(tmp_path):
    path = workbook(tmp_path / "t.xlsx", {"T": [["a", "b", "c"], [1, 2, 3], [4, 5, 6]]})
    temporary = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb'<dimension ref="[^"]*"\s*/>', b'<dimension ref="A1"/>', data)  # a writer's mistake
            target.writestr(item, data)
    temporary.replace(path)

    table = read(path)

    assert table.column_names == ["a", "b", "c"] and table.num_rows == 2


# --------------------------------------------------------------------------- the header cells


def test_header_cells_lose_extra_spaces_and_line_breaks_and_numbers_become_text(tmp_path):
    path = workbook(tmp_path / "t.xlsx", {"T": [["  order\nid ", "unit   price", 2026, 1.5], [1, 2, 3, 4]]})

    assert read(path).column_names == ["order id", "unit price", "2026", "1.5"]


def test_an_empty_header_cell_names_the_cell_the_file_and_the_sheet(tmp_path):
    path = workbook(tmp_path / "t.xlsx", {"Data": [["a", None, "c"], [1, 2, 3]]})

    with pytest.raises(ConfigError) as error:
        read(path)

    assert str(error.value) == (f"{path} (sheet 'Data'): the header cell B1 is empty; every column up to the last "
                                f"one that has data needs a name")


def test_an_empty_header_cell_above_data_is_an_error_on_the_header_row_it_was_read_from(tmp_path):
    path = workbook(tmp_path / "t.xlsx", {"Data": [["title"], ["a", "b", None], [1, 2, 3]]})

    with pytest.raises(ConfigError, match=r"the header cell C2 is empty"):
        read(path, header_row=2)


@pytest.mark.parametrize("header", [["id", "name", "id"], ["id", "name", "ID"], ["id", "name", " Id "]])
def test_a_repeated_header_names_both_cells(tmp_path, header):
    path = workbook(tmp_path / "t.xlsx", {"Data": [header, [1, "a", 2]]})

    with pytest.raises(ConfigError, match=rf"\(sheet 'Data'\): the header '{header[2].strip()}' is in both A1 and C1"):
        read(path)


# --------------------------------------------------------------------------- files that cannot be read


def test_an_old_xls_workbook_is_refused_with_what_to_do(tmp_path):
    path = tmp_path / "old.xls"
    path.write_bytes(b"\xd0\xcf\x11\xe0 an old workbook")

    with pytest.raises(ConfigError) as error:
        read(path)

    assert str(error.value) == (f"{path} is an old .xls workbook, which is not supported. Open it in Excel and "
                                f"save it as .xlsx")


def test_a_file_that_is_not_a_workbook_names_the_file(tmp_path):
    path = tmp_path / "fake.xlsx"
    path.write_bytes(b"this is not a zip file")

    with pytest.raises(ConfigError, match=rf"{re.escape(str(path))} is not a readable .xlsx workbook"):
        read(path)


def test_a_file_with_another_extension_is_refused_by_name(tmp_path):
    path = tmp_path / "sales.csv"
    path.write_text("a,b\n1,2\n")

    with pytest.raises(ConfigError, match=rf"{re.escape(str(path))} cannot be read: .*does not support .csv"):
        read(path)


def test_a_workbook_that_does_not_exist_is_a_file_not_found_error(tmp_path):
    with pytest.raises(FileNotFoundError, match="Excel file not found: .*nope.xlsx"):
        read(tmp_path / "nope.xlsx")


def test_a_workbook_made_to_exhaust_memory_is_refused(tmp_path):
    bomb = ('<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaaaaaaaa"><!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;&a;&a;">'
            '<!ENTITY c "&b;&b;&b;&b;&b;&b;&b;&b;&b;&b;">]><worksheet '
            'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1">'
            '<c r="A1" t="inlineStr"><is><t>&c;</t></is></c></row></sheetData></worksheet>')
    path = workbook(tmp_path / "bomb.xlsx", {"T": [["a"]]})
    temporary = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            target.writestr(item, bomb.encode() if item.filename == "xl/worksheets/sheet1.xml"
                            else source.read(item.filename))
    temporary.replace(path)

    with pytest.raises(ConfigError, match="is not a readable .xlsx workbook"):
        read(path)


def test_without_openpyxl_the_error_is_the_install_command(sales, monkeypatch):
    monkeypatch.setitem(sys.modules, "openpyxl", None)  # makes `import openpyxl` fail, as it does without the extra

    with pytest.raises(EngineNotFound) as error:
        read(sales)

    assert 'pip install "local-data-platform[excel]"' in str(error.value) and INSTALL_HINT in str(error.value)


def test_without_defusedxml_there_is_one_warning(sales, monkeypatch, caplog):
    monkeypatch.setattr(openpyxl, "DEFUSEDXML", False)
    monkeypatch.setattr(excel_module, "_warned_about_defusedxml", False)
    caplog.set_level(logging.WARNING, logger="local_data_platform.format.excel")

    read(sales, sheet="Menu")
    read(sales, sheet="Menu")

    warnings = [record.getMessage() for record in caplog.records
                if "defusedxml is not installed" in record.getMessage()]
    assert len(warnings) == 1 and INSTALL_HINT in warnings[0]


# --------------------------------------------------------------------------- the source object


@pytest.mark.parametrize("options, message", [
    ({"sheet": True}, "'sheet' must be a sheet name or a position from 0"),
    ({"sheet": -1}, "'sheet' must be a sheet name or a position from 0"),
    ({"sheet": 1.5}, "'sheet' must be a sheet name or a position from 0"),
    ({"sheet": " "}, "'sheet' must be a sheet name or a position from 0"),
    ({"header_row": 0}, "'header_row' must be a whole number from 1"),
    ({"header_row": "1"}, "'header_row' must be a whole number from 1"),
    ({"header_row": True}, "'header_row' must be a whole number from 1"),
])
def test_a_sheet_or_header_row_that_makes_no_sense_is_refused_when_the_source_is_built(options, message):
    with pytest.raises(ConfigError, match=message):
        Excel("sales", "sales.xlsx", **options)


def test_a_path_is_needed_and_must_be_local():
    with pytest.raises(ConfigError, match="needs a path"):
        Excel("sales", "")
    with pytest.raises(ConfigError, match="is not a local file"):
        Excel("sales", "s3://bucket/sales.xlsx")


def test_a_relative_path_resolves_against_the_base_folder_and_the_source_knows_its_format(tmp_path):
    source = Excel("sales", "data/sales.xlsx", sheet=0, base_dir=tmp_path)

    assert source.path == tmp_path / "data" / "sales.xlsx" and source.format == "XLSX"
    assert source.sheet == 0 and source.header_row == 1
    assert repr(source) == f"Excel(name='sales', path='{tmp_path / 'data' / 'sales.xlsx'}', sheet=0, header_row=1)"


def test_the_source_cannot_be_written_to(sales):
    with pytest.raises(Exception, match="cannot be written"):
        Excel("sales", sales).put(pa.table({"a": [1]}))


def test_nothing_is_changed_in_the_workbook_and_it_is_not_left_open(sales):
    before = sales.read_bytes()

    read(sales, sheet="Sales", header_row=2)

    assert sales.read_bytes() == before
    sales.unlink()  # a handle left open would make this fail on Windows, and it shows the file is released


# --------------------------------------------------------------------------- the config, the spec and the routes


def config_for(folder: Path, source: dict, target: dict) -> Config:
    return Config.from_dict({"identifier": "sales", "metadata": {
        "source": {"name": "sales", "format": "XLSX", "path": "sales.xlsx", **source},
        "target": target}}, base_dir=folder)


def csv_target(path: str = "out/sales.csv") -> dict:
    return {"name": "sales", "format": "CSV", "path": path}


def iceberg_target(folder: Path, **options) -> dict:
    return {"name": "sales", "format": "ICEBERG", "catalog": {"identifier": "ns", "warehouse_path": str(folder / "wh")},
            **options}


def test_the_routes_are_registered_and_xlsx_is_a_known_format():
    assert "XLSX" in FORMATS and SupportedFormat.XLSX.value == "XLSX"
    assert get_pipeline_class(*make_route("XLSX", "CSV")).__name__ == "ExcelToCSV"
    assert get_pipeline_class(*make_route("xlsx", "ICEBERG")).__name__ == "ExcelToIceberg"


def test_excel_from_config_reads_the_block_and_warns_about_keys_it_ignores(tmp_path, caplog):
    config = config_for(tmp_path, {"sheet": "Sales", "header_row": 3, "engine": "DUCKDB"}, csv_target())
    caplog.set_level(logging.WARNING)

    source = excel_from_config(config)

    assert source.sheet == "Sales" and source.header_row == 3 and source.path == tmp_path / "sales.xlsx"
    assert any("keys ['engine'] are not used" in record.getMessage() for record in caplog.records)
    with pytest.raises(ConfigError, match="expected 'XLSX'"):
        excel_from_config(config, "target")


def test_excel_to_csv_writes_the_sheet_as_a_csv_file(tmp_path, sales):
    config = config_for(tmp_path, {"sheet": "Sales", "header_row": 2}, csv_target())

    result = create_pipeline(config).run()

    assert result.rows_written == 2
    assert (tmp_path / "out" / "sales.csv").read_text().splitlines() == [
        '"day","dish","orders","revenue"',
        '2026-07-01 00:00:00.000000,"pad thai",12,143.5',
        '2026-07-02 00:00:00.000000,"green curry",9,117']


def test_the_csv_is_replaced_atomically_and_a_failed_read_leaves_it_as_it_was(tmp_path, sales):
    target = tmp_path / "out" / "sales.csv"
    target.parent.mkdir()
    target.write_text("old,contents\n1,2\n")

    with pytest.raises(ConfigError, match="there is no sheet named 'Nope'"):
        create_pipeline(config_for(tmp_path, {"sheet": "Nope"}, csv_target())).run()
    assert target.read_text() == "old,contents\n1,2\n"

    create_pipeline(config_for(tmp_path, {"sheet": "Menu"}, csv_target())).run()
    assert target.read_text().splitlines()[0] == '"dish","price"'
    assert [path.name for path in target.parent.iterdir()] == ["sales.csv"]  # no temporary file is left behind


def test_excel_to_iceberg_keeps_the_types_and_the_rows(tmp_path, sales):
    config = config_for(tmp_path, {"sheet": "Sales", "header_row": 2},
                        iceberg_target(tmp_path, write_mode="overwrite"))

    create_pipeline(config).run()
    table = Iceberg("sales", {"identifier": "ns", "warehouse_path": str(tmp_path / "wh")}).get()

    assert {field.name: str(field.type) for field in table.schema} == {
        "day": "timestamp[us]", "dish": "string", "orders": "int64", "revenue": "double"}
    assert table.sort_by("day").column("orders").to_pylist() == [12, 9]
    again = create_pipeline(config).run()  # overwrite is idempotent
    assert again.write_result.rows_after == 2


def test_excel_to_iceberg_upserts_on_a_key(tmp_path, sales):
    first = config_for(tmp_path, {"sheet": "Menu"}, iceberg_target(tmp_path, write_mode="upsert", join_cols=["dish"]))
    create_pipeline(first).run()
    workbook(tmp_path / "sales.xlsx", {"Menu": [["dish", "price"], ["pad thai", 12.0], ["tom yum", 14.5]]})

    result = create_pipeline(first).run()

    table = Iceberg("sales", {"identifier": "ns", "warehouse_path": str(tmp_path / "wh")}).get()
    assert sorted(table.to_pylist(), key=lambda row: row["dish"]) == [
        {"dish": "green curry", "price": 13.0}, {"dish": "pad thai", "price": 12.0},
        {"dish": "tom yum", "price": 14.5}]
    assert result.rows_written == 2  # one changed and one new


def test_quality_checks_block_a_bad_sheet_before_anything_is_written(tmp_path, sales):
    config = config_for(tmp_path, {"sheet": "Sales", "header_row": 2}, iceberg_target(tmp_path))
    config.metadata["quality"] = {"on_failure": "fail", "checks": [{"check": "range", "column": "orders", "min": 10}]}

    with pytest.raises(DataQualityError):
        create_pipeline(config).run()

    assert not Iceberg("sales", {"identifier": "ns", "warehouse_path": str(tmp_path / "wh")}).exists()


def test_run_events_name_the_workbook(tmp_path, sales):
    assert dataset_ref(Excel("sales", sales)) == {"namespace": "file", "name": sales.resolve().as_posix(),
                                                  "format": "XLSX"}


def test_ldp_run_turns_a_workbook_into_a_csv(tmp_path, sales, capsys):
    config = tmp_path / "sales.json"
    config.write_text(json.dumps({"identifier": "sales", "metadata": {
        "source": {"name": "sales", "format": "XLSX", "path": "sales.xlsx", "sheet": "Sales", "header_row": 2},
        "target": csv_target("drive/sales.csv")}}))

    assert main(["run", str(config)]) == 0

    assert "sales: read 2 rows, wrote 2 rows" in capsys.readouterr().out
    assert (tmp_path / "drive" / "sales.csv").is_file()


def test_ldp_run_without_the_extra_prints_the_install_command(tmp_path, sales, capsys, monkeypatch):
    config = tmp_path / "sales.json"
    config.write_text(json.dumps({"identifier": "sales", "metadata": {
        "source": {"name": "sales", "format": "XLSX", "path": "sales.xlsx"}, "target": csv_target()}}))
    monkeypatch.setitem(sys.modules, "openpyxl", None)

    assert main(["run", str(config)]) == 1

    assert ('ldp: error: An XLSX source needs the openpyxl package. Install it with: '
            'pip install "local-data-platform[excel]"') in capsys.readouterr().err


def test_ldp_run_on_a_bad_sheet_names_the_file_and_the_sheet(tmp_path, sales, capsys):
    config = tmp_path / "sales.json"
    config.write_text(json.dumps({"identifier": "sales", "metadata": {
        "source": {"name": "sales", "format": "XLSX", "path": "sales.xlsx", "sheet": "Nope"},
        "target": csv_target()}}))

    assert main(["run", str(config)]) == 1

    err = capsys.readouterr().err
    assert err.startswith(f"ldp: error: {tmp_path / 'sales.xlsx'}: there is no sheet named 'Nope'; the sheets are")


# --------------------------------------------------------------------------- the spec, the JSON schema, ldp plan


def spec(source: dict | None = None, target: dict | None = None) -> dict:
    return {"apiVersion": "ldp/v1", "identifier": "sales", "metadata": {
        "source": {"name": "sales", "format": "XLSX", "path": "sales.xlsx", **(source or {})},
        "target": target or csv_target()}}


@pytest.mark.parametrize("source", [{}, {"sheet": "Sales"}, {"sheet": 0}, {"sheet": 3, "header_row": 5},
                                    {"header_row": 1}])
def test_validate_spec_accepts_an_xlsx_source(source):
    assert validate_spec(spec(source)) == []
    assert validate_spec(spec({**source, "format": "xlsx"})) == []


@pytest.mark.parametrize("source, problem", [
    ({"sheet": True}, "metadata.source.sheet: must be a sheet name or a position from 0, got True"),
    ({"sheet": -1}, "metadata.source.sheet: must be a sheet name or a position from 0, got -1"),
    ({"sheet": ""}, "metadata.source.sheet: must be a sheet name or a position from 0, got ''"),
    ({"header_row": 0}, "metadata.source.header_row: must be a whole number from 1, got 0"),
    ({"header_row": "1"}, "metadata.source.header_row: must be a whole number from 1, got '1'"),
    ({"path": "old.xls"}, "metadata.source.path: is an old .xls workbook, which is not supported"),
    ({"path": None}, "metadata.source.path: is required"),
])
def test_validate_spec_names_what_is_wrong_in_an_xlsx_source(source, problem):
    assert any(message.startswith(problem) for message in map(str, validate_spec(spec(source)))), source


def test_an_xlsx_block_cannot_be_a_target():
    target = {"name": "out", "format": "XLSX", "path": "o.xlsx"}
    problems = [str(problem) for problem in validate_spec(spec(target=target))]

    assert "metadata.target.format: XLSX is a source format only; write to CSV or ICEBERG instead" in problems


def test_the_json_schema_accepts_and_rejects_the_same_blocks():
    jsonschema = pytest.importorskip("jsonschema")
    validator = jsonschema.Draft202012Validator(json_schema())

    for good in (spec(), spec({"sheet": "Sales", "header_row": 2}), spec({"sheet": 0})):
        assert not list(validator.iter_errors(good)), good
    for bad in (spec({"sheet": -1}), spec({"sheet": True}), spec({"header_row": 0}), spec({"header_row": "2"}),
                spec(target={"name": "o", "format": "XLSX", "path": "o.xlsx"})):
        assert list(validator.iter_errors(bad)), bad


def test_ldp_plan_shows_the_route_and_reports_a_missing_workbook(tmp_path, sales):
    good = tmp_path / "sales.json"
    good.write_text(json.dumps({"identifier": "sales", "metadata": {
        "source": {"name": "sales", "format": "XLSX", "path": "sales.xlsx"}, "target": csv_target()}}))
    missing = tmp_path / "missing.json"
    missing.write_text(json.dumps({"identifier": "missing", "metadata": {
        "source": {"name": "sales", "format": "XLSX", "path": "nope.xlsx"}, "target": csv_target()}}))

    result = plan(good)
    gone = plan(missing)

    assert result["valid"] is True and result["route"] == "XLSX -> CSV" and result["pipeline"] == "ExcelToCSV"
    assert gone["valid"] is False
    assert gone["errors"] == [f"metadata.source.path: the Excel file {tmp_path / 'nope.xlsx'} does not exist"]
