"""An Excel workbook (``.xlsx``) as a source: one sheet read into a ``pyarrow.Table``.

``openpyxl`` is an optional dependency (the ``excel`` extra). It is imported when a sheet is read, never when
this module is imported::

    pip install "local-data-platform[excel]"

The extra also installs ``defusedxml``, which ``openpyxl`` uses on its own to refuse a spreadsheet built to
exhaust memory. A workbook is opened read-only and closed before :meth:`Excel.get` returns.

The first row of the table is the header row (``header_row``, 1 by default), and the rows above it are skipped.
Each column becomes one Arrow type:

| The cells hold | The column is |
|---|---|
| only whole numbers | ``int64`` |
| numbers, some not whole | ``float64`` |
| only dates and datetimes | ``timestamp[us]``, with no time zone |
| only times of day | ``time64[us]`` |
| only TRUE and FALSE | ``bool`` |
| only text, or nothing at all | ``string`` |
| more than one of those | ``string``, and a warning names the column |

Empty cells are nulls. A formula cell gives the value Excel saved with it, which is empty when the file was
written by a program that never calculated it. Empty rows at the end of the sheet and empty columns at its
right edge are dropped.
"""

from __future__ import annotations

import datetime as dt
import os
import zipfile
from typing import Any

import pyarrow as pa

from local_data_platform.exceptions import ConfigError, EngineNotFound
from local_data_platform.format import Format
from local_data_platform.logger import get_logger

logger = get_logger(__name__)

INSTALL_HINT = 'pip install "local-data-platform[excel]"'
EXTENSIONS = (".xlsx", ".xlsm", ".xltx", ".xltm")
_INT64_LIMIT = 2**63
_warned_about_defusedxml = False


def _import_openpyxl() -> Any:
    """Import ``openpyxl`` or raise :class:`EngineNotFound` with an install hint."""
    global _warned_about_defusedxml
    try:
        import openpyxl
    except ImportError as error:
        raise EngineNotFound(f"An XLSX source needs the openpyxl package. Install it with: {INSTALL_HINT}") from error
    if not getattr(openpyxl, "DEFUSEDXML", True) and not _warned_about_defusedxml:
        _warned_about_defusedxml = True
        logger.warning("defusedxml is not installed, so a crafted spreadsheet could use all of the memory when it is "
                       "read. Install it with: %s", INSTALL_HINT)
    return openpyxl


class Excel(Format):
    """One sheet of an ``.xlsx`` workbook on the local disk, read-only.

    Args:
        name: Dataset name.
        path: The workbook. Relative paths resolve against ``base_dir`` (or the cwd).
        sheet: The sheet: its name, or its position counted from 0. The first sheet by default.
        header_row: The row that holds the column names, counted from 1. Rows above it are skipped.
        base_dir: Folder that relative paths resolve against, usually the config file's folder.
        **options: Extra options kept on ``self.options``.

    Attributes:
        path: The absolute path of the workbook.
        sheet: The sheet as given, or ``None`` for the first one.
        header_row: The header row.

    Raises:
        ConfigError: If ``path`` is empty or not a local file, ``sheet`` is not a name or a position, or
            ``header_row`` is not a whole number of at least 1.
    """

    def __init__(self, name: str, path: str | os.PathLike, sheet: str | int | None = None, header_row: int = 1,
                 base_dir: str | os.PathLike | None = None, **options):
        if not path:
            raise ConfigError(f"Excel dataset {name!r} needs a path")
        if "://" in str(path):
            raise ConfigError(f"Excel dataset {name!r}: {path} is not a local file; an XLSX source reads a workbook "
                              f"from the local disk")
        if sheet is not None and not (isinstance(sheet, str) and sheet.strip()) and not (
                isinstance(sheet, int) and not isinstance(sheet, bool) and sheet >= 0):
            raise ConfigError(f"Excel dataset {name!r}: 'sheet' must be a sheet name or a position from 0, "
                              f"got {sheet!r}")
        if not isinstance(header_row, int) or isinstance(header_row, bool) or header_row < 1:
            raise ConfigError(f"Excel dataset {name!r}: 'header_row' must be a whole number from 1, "
                              f"got {header_row!r}")
        super().__init__(name, path=path, format="XLSX", base_dir=base_dir, **options)
        self.sheet = sheet
        self.header_row = header_row

    def get(self) -> pa.Table:
        """Read the sheet.

        Returns:
            The data rows below the header row as a ``pyarrow.Table``, one column for each header cell.

        Raises:
            EngineNotFound: If ``openpyxl`` is not installed.
            FileNotFoundError: If the workbook does not exist.
            ConfigError: If the file is an old ``.xls`` workbook or is not a workbook at all, the sheet does
                not exist, or a header cell is empty or repeated. The message names the file and the sheet.
        """
        openpyxl = _import_openpyxl()
        if not self.path.is_file():
            raise FileNotFoundError(f"Excel file not found: {self.path}")
        if self.path.suffix.lower() == ".xls":
            raise ConfigError(f"{self.path} is an old .xls workbook, which is not supported. Open it in Excel and "
                              f"save it as .xlsx")
        try:
            workbook = openpyxl.load_workbook(self.path, read_only=True, data_only=True)
        except zipfile.BadZipFile as error:
            raise ConfigError(f"{self.path} is not a readable .xlsx workbook. A file that is password protected, or "
                              f"that has another format and an .xlsx name, cannot be read") from error
        except openpyxl.utils.exceptions.InvalidFileException as error:
            raise ConfigError(f"{self.path} cannot be read: {error}") from error
        except (KeyError, ValueError) as error:  # what openpyxl raises for a workbook whose XML is invalid
            raise ConfigError(f"{self.path} is not a readable .xlsx workbook: {_first_line(error)}") from error
        where = str(self.path)
        try:
            worksheet = self._select(workbook)
            where = f"{self.path} (sheet {worksheet.title!r})"
            worksheet.reset_dimensions()  # some writers record a wrong size, and read-only mode trusts it
            rows = [tuple(row) for row in worksheet.iter_rows(min_row=self.header_row, values_only=True)]
        except (zipfile.BadZipFile, KeyError, ValueError, SyntaxError) as error:
            raise ConfigError(f"{where}: the sheet cannot be read: {_first_line(error)}") from error
        finally:
            workbook.close()
        table = _table(rows, self.header_row, where)
        logger.info("Read %d rows and %d columns from Excel %s", table.num_rows, table.num_columns, where)
        return table

    def _select(self, workbook: Any) -> Any:
        """The worksheet to read: by position, by name (Excel ignores case in names), or the first."""
        sheets = workbook.worksheets
        names = [sheet.title for sheet in sheets]
        if not sheets:
            raise ConfigError(f"{self.path} has no worksheets")
        if self.sheet is None:
            return sheets[0]
        if isinstance(self.sheet, int):
            if self.sheet >= len(sheets):
                raise ConfigError(f"{self.path}: there is no sheet at position {self.sheet}; the workbook has "
                                  f"{len(sheets)}: {names}")
            return sheets[self.sheet]
        for sheet in sheets:
            if sheet.title == self.sheet or sheet.title.casefold() == self.sheet.strip().casefold():
                return sheet
        raise ConfigError(f"{self.path}: there is no sheet named {self.sheet!r}; the sheets are {names}")

    def __repr__(self) -> str:
        return f"Excel(name={self.name!r}, path={str(self.path)!r}, sheet={self.sheet!r}, header_row={self.header_row})"


# --------------------------------------------------------------------------- from rows to a table


def _table(rows: list[tuple], header_row: int, where: str) -> pa.Table:
    """Build the table from the rows read from ``header_row`` down: a header row, then data rows."""
    if not rows:
        raise ConfigError(f"{where}: the sheet has no row {header_row}, so there is no header to read")
    width = max(len(row) for row in rows)
    rows = [row + (None,) * (width - len(row)) for row in rows]
    while len(rows) > 1 and all(_is_empty(cell) for cell in rows[-1]):
        rows.pop()  # empty rows at the end, often left behind by formatting
    header, data = rows[0], rows[1:]
    columns = [[row[index] for row in data] for index in range(width)]
    last = max((index for index in range(width) if not _is_blank(header[index])
                or any(not _is_empty(cell) for cell in columns[index])), default=-1)
    if last < 0:
        raise ConfigError(f"{where}: row {header_row}, the header row, is empty")
    names = _names(header[:last + 1], header_row, where)
    return pa.table({name: _array(values, name, where) for name, values in zip(names, columns)})


def _names(header: tuple, header_row: int, where: str) -> list[str]:
    from openpyxl.utils import get_column_letter

    names: list[str] = []
    seen: dict[str, str] = {}
    for index, cell in enumerate(header):
        position = f"{get_column_letter(index + 1)}{header_row}"
        if _is_blank(cell):
            raise ConfigError(f"{where}: the header cell {position} is empty; every column up to the last one "
                              f"that has data needs a name")
        name = " ".join(str(cell).split())  # one space for any run of spaces and line breaks
        if name.casefold() in seen:
            raise ConfigError(f"{where}: the header {name!r} is in both {seen[name.casefold()]} and {position}; "
                              f"every column needs its own name, whatever the case")
        seen[name.casefold()] = position
        names.append(name)
    return names


def _first_line(error: BaseException) -> str:
    text = str(error).strip()
    return text.splitlines()[0] if text else type(error).__name__


def _is_empty(cell: Any) -> bool:
    return cell is None or (isinstance(cell, str) and cell == "")


def _is_blank(cell: Any) -> bool:
    return cell is None or (isinstance(cell, str) and not cell.strip())


def _kind(cell: Any) -> str:
    if isinstance(cell, bool):
        return "bool"
    if isinstance(cell, (int, float)):
        return "number"
    if isinstance(cell, str):
        return "text"
    if isinstance(cell, (dt.datetime, dt.date)):
        return "date"
    if isinstance(cell, dt.time):
        return "time"
    return "other"


def _array(values: list[Any], name: str, where: str) -> pa.Array:
    """One column: the Arrow type that fits all of its cells, or text when they do not agree."""
    cells = [None if _is_empty(value) else value for value in values]
    kinds = {_kind(cell) for cell in cells if cell is not None}
    if not kinds:
        return pa.array([None] * len(cells), pa.string())
    if len(kinds) > 1:
        logger.warning("%s: column %r holds %s, so it is read as text", where, name, " and ".join(sorted(kinds)))
        return pa.array([_text(cell) for cell in cells], pa.string())
    (kind,) = kinds
    if kind == "bool":
        return pa.array(cells, pa.bool_())
    if kind == "text":
        return pa.array(cells, pa.string())
    if kind == "date":
        return pa.array([None if cell is None else _as_datetime(cell) for cell in cells], pa.timestamp("us"))
    if kind == "time":
        return pa.array(cells, pa.time64("us"))
    if kind == "number":
        numbers = [cell for cell in cells if cell is not None]
        if all(_is_whole(number) for number in numbers):
            return pa.array([None if cell is None else int(cell) for cell in cells], pa.int64())
        try:
            return pa.array([None if cell is None else float(cell) for cell in cells], pa.float64())
        except OverflowError:  # an integer too large for any number type, from a file made to break readers
            logger.warning("%s: column %r holds numbers too large to keep, so it is read as text", where, name)
    return pa.array([_text(cell) for cell in cells], pa.string())


def _is_whole(number: int | float) -> bool:
    """Whether ``number`` is a whole number that fits in an ``int64``."""
    return abs(number) < _INT64_LIMIT and (isinstance(number, int) or number.is_integer())


def _as_datetime(value: dt.date) -> dt.datetime:
    return value if isinstance(value, dt.datetime) else dt.datetime.combine(value, dt.time())


def _text(value: Any) -> str | None:
    """A cell as the text of a column that mixes types."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, dt.datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    return str(value)


__all__ = ["EXTENSIONS", "Excel", "INSTALL_HINT"]
